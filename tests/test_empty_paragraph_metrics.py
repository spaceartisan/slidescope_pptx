"""Regression: DrawingML endParaRPr controls the height of empty paragraphs."""
import io
import unittest
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx_parser import parse_pptx


class EndParagraphFontMetricsTests(unittest.TestCase):
    def test_empty_paragraph_uses_own_end_run_size(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(914400, 914400, 7315200, 3657600)
        shape.text_frame.paragraphs[0].text = 'Visible introduction'
        empty = shape.text_frame.add_paragraph()
        end = OxmlElement('a:endParaRPr')
        end.set('sz', '1900')  # 19pt -> 25.333px (not the theme default)
        empty._p.append(end)
        shape.text_frame.add_paragraph().text = 'Following paragraph'
        out = io.BytesIO()
        deck.save(out)
        model, _ = parse_pptx(out.getvalue(), 'font-metrics.pptx')
        shapes = [s for s in model['slides'][0]['shapes'] if s['source'] == 'slide'
                  and s.get('text') and len(s['text']['paragraphs']) == 3]
        self.assertEqual(len(shapes), 1)
        blank = shapes[0]['text']['paragraphs'][1]
        self.assertEqual(blank['text'], '')
        self.assertAlmostEqual(blank['style']['size'], 25.333, places=2)


if __name__ == '__main__':
    unittest.main()
