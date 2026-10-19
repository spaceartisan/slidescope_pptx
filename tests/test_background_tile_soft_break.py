"""Regression tests for native tile-scale and DrawingML soft breaks.

Optional real deck: SLIDESCOPE_REGRESSION_PPTX=path/to/deck.pptx
"""
from pathlib import Path
from io import BytesIO
import os
import unittest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx_parser import parse_pptx


class SoftBreakTests(unittest.TestCase):
    def test_soft_break_preserved_in_renderable_runs(self):
        prs=Presentation()
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
        para=box.text_frame.paragraphs[0]
        para.add_run().text='Title'
        para.add_line_break()
        para.add_run().text='Subtitle'
        data=BytesIO();prs.save(data)
        model,_=parse_pptx(data.getvalue())
        body=next(s for s in model['slides'][0]['shapes'] if s['source']=='slide' and s.get('text'))
        p=body['text']['paragraphs'][0]
        self.assertEqual(p['text'],'Title\vSubtitle')
        self.assertEqual(''.join(r['text'] for r in p['runs']),'Title\nSubtitle')
        self.assertEqual([r['text'] for r in p['runs']],['Title','\n','Subtitle'])

    def test_title_slide_background_preserves_ooxml_tile(self):
        source=os.environ.get('SLIDESCOPE_REGRESSION_PPTX')
        if not source:
            self.skipTest('Needs the original OOXML regression deck')
        model,_=parse_pptx(Path(source).read_bytes())
        self.assertEqual(model['count'],12)
        for num in (1,9):
            tile=model['slides'][num-1]['backgroundImageTile']
            self.assertIsNotNone(tile)
            self.assertAlmostEqual(tile['w'],1174,delta=5)
            self.assertAlmostEqual(tile['h'],880,delta=5)
        for num in range(2,13):
            if num!=9:
                self.assertIsNone(model['slides'][num-1]['backgroundImageTile'])
        title=next(s for s in model['slides'][8]['shapes'] if s['source']=='slide'
                   and s.get('text') and 'Practice Examples' in s['text']['paragraphs'][0]['text'])
        rendered=''.join(r['text'] for r in title['text']['paragraphs'][0]['runs'])
        self.assertIn('\n',rendered)

if __name__=='__main__':
    unittest.main()
