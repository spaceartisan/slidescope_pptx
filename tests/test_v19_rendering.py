"""v1.9: inherited initial paragraph leading and basic OOXML auto-numbering."""
import io
import unittest
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches
from pptx_parser import parse_pptx

class V19Tests(unittest.TestCase):
    def test_bullet_auto_numbering_preserves_format(self):
        deck=Presentation()
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2))
        para=shape.text_frame.paragraphs[0]
        para.text='Example'
        ns='http://schemas.openxmlformats.org/drawingml/2006/main'
        ppr=etree.Element('{%s}pPr' % ns)
        el=etree.SubElement(ppr,'{%s}buAutoNum'%ns)
        el.set('type','romanLcPeriod')
        el.set('startAt','3')
        para._p.insert(0,ppr)
        data=io.BytesIO();deck.save(data)
        model,_=parse_pptx(data.getvalue())
        paragraph=next(s for s in model['slides'][0]['shapes'] if s.get('text'))['text']['paragraphs'][0]
        self.assertTrue(paragraph['autoBullet'])
        self.assertIsNone(paragraph['bullet'])
        self.assertEqual(paragraph['numbering'],{'type':'romanLcPeriod','startAt':3})

    def test_browser_honors_non_wrapping_text(self):
        js=(Path(__file__).resolve().parents[1]/'static/viewer.js').read_text()
        self.assertIn("textModel.wrap === false",js)
        self.assertIn("para.style.overflowWrap='normal'",js)
        self.assertIn('powerPointAutoNumber',js)

    def test_first_paragraph_leading_keeps_later_text_in_place(self):
        js=(Path(__file__).resolve().parents[1]/'static/viewer.js').read_text()
        self.assertIn('const transferFirstLeading = index===0 && opts.isPlaceholder',js)
        self.assertIn('textModel.autoSize == null && p.spaceBefore == null && before>0',js)
        self.assertIn('after+(transferFirstLeading?before:0)',js)

if __name__=='__main__': unittest.main()
