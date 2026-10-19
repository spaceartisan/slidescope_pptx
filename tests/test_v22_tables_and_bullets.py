"""General DrawingML regression cases, independent of the reference decks."""
import unittest
from io import BytesIO
from lxml import etree
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.xmlchemy import OxmlElement

from pptx_parser import parse_pptx, color_node, _table_regions

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


class DrawingMLTintTests(unittest.TestCase):
    def test_tint_endpoints_and_linear_light_mix(self):
        for value, expected in [(0, '#FFFFFF'), (100000, '#17365D'), (40000, '#CCCED2')]:
            node=etree.fromstring(f'<a:srgbClr xmlns:a="{A}" val="17365D"><a:tint val="{value}"/></a:srgbClr>')
            self.assertEqual(color_node(node,{}),expected)


class TableFormattingTests(unittest.TestCase):
    def model(self, direct=False, merged=False):
        prs=Presentation()
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_table(4,2, Inches(1),Inches(1),Inches(6),Inches(3))
        table=shape.table
        table.first_row=True
        table.horz_banding=True
        for r,row in enumerate(table.rows):
            for c,cell in enumerate(row.cells):
                cell.text=f'Cell {r},{c}'
        if merged: table.cell(0,0).merge(table.cell(0,1))
        # Replace the generated style part, as a producer might customize it.
        style_id=table._tbl.tblPr.find('{'+A+'}tableStyleId').text
        styles=f'''<a:tblStyleLst xmlns:a="{A}" def="{style_id}"><a:tblStyle styleId="{style_id}" styleName="Synthetic">
        <a:wholeTbl><a:tcTxStyle><a:srgbClr val="223344"/></a:tcTxStyle><a:tcStyle><a:fill><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:fill></a:tcStyle></a:wholeTbl>
        <a:band1H><a:tcStyle><a:fill><a:solidFill><a:srgbClr val="DDDDDD"/></a:solidFill></a:fill></a:tcStyle></a:band1H>
        <a:firstRow><a:tcTxStyle b="on"><a:srgbClr val="FFFFFF"/></a:tcTxStyle><a:tcStyle><a:fill><a:solidFill><a:srgbClr val="123456"/></a:solidFill></a:fill></a:tcStyle></a:firstRow>
        </a:tblStyle></a:tblStyleLst>'''
        for part in prs.part.package.iter_parts():
            if str(part.partname).endswith('/tableStyles.xml'): part._blob=styles.encode()
        if direct:
            cell=table.cell(0,0)
            from pptx.dml.color import RGBColor
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(255,0,0)
            cell.text_frame.paragraphs[0].runs[0].font.bold=False
            cell.text_frame.paragraphs[0].runs[0].font.color.rgb=RGBColor(0,255,0)
            cell.margin_left=Inches(.2)
            ln=OxmlElement('a:lnL');ln.set('w','25400')
            solid=OxmlElement('a:solidFill');col=OxmlElement('a:srgbClr');col.set('val','ABCDEF');solid.append(col);ln.append(solid)
            cell._tc.get_or_add_tcPr().append(ln)
        stream=BytesIO();prs.save(stream)
        model,_=parse_pptx(stream.getvalue())
        return next(o['table'] for o in model['slides'][0]['shapes'] if o['kind']=='table')

    def test_header_and_bands_use_package_style(self):
        table=self.model()
        self.assertEqual([r[0]['fill'] for r in table['cells']],['#123456','#DDDDDD','#FFFFFF','#DDDDDD'])
        run=table['cells'][0][0]['text']['paragraphs'][0]['runs'][0]
        self.assertTrue(run['bold']);self.assertEqual(run['color'],'#FFFFFF')
        self.assertEqual(table['cells'][1][0]['text']['paragraphs'][0]['runs'][0]['color'],'#223344')

    def test_explicit_cell_and_run_formatting_override_style(self):
        cell=self.model(direct=True)['cells'][0][0]
        self.assertEqual(cell['fill'],'#FF0000')
        run=cell['text']['paragraphs'][0]['runs'][0]
        self.assertFalse(run['bold']);self.assertEqual(run['color'],'#00FF00')
        self.assertAlmostEqual(cell['margin'][3],19.2)
        self.assertAlmostEqual(cell['borders']['left']['width'],2.667,places=3)
        self.assertEqual(cell['borders']['left']['color'],'#ABCDEF')

    def test_merged_cell_span_survives_style_resolution(self):
        cells=self.model(merged=True)['cells']
        self.assertEqual(cells[0][0]['spanWidth'],2)
        self.assertTrue(cells[0][1]['spanned'])
        self.assertEqual(cells[0][0]['fill'],'#123456')

    def test_corner_style_requires_both_enabled_edges(self):
        props=etree.fromstring(f'<a:tblPr xmlns:a="{A}" firstRow="1" firstCol="1"/>')
        self.assertEqual(_table_regions(props,0,0,3,3)[-1],'nwCell')
        props.set('firstCol','0')
        self.assertNotIn('nwCell',_table_regions(props,0,0,3,3))


class BulletFontTests(unittest.TestCase):
    def test_local_bullet_font_and_follow_text_override(self):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[1])
        shape=slide.placeholders[1];shape.text='First'
        p=shape.text_frame.paragraphs[0]
        pp=p._p.get_or_add_pPr()
        for tag, attrs in [('a:buFont',{'typeface':'Wingdings'}),('a:buChar',{'char':'§'})]:
            el=OxmlElement(tag)
            for k,v in attrs.items():el.set(k,v)
            pp.append(el)
        stream=BytesIO();prs.save(stream)
        model,_=parse_pptx(stream.getvalue())
        body=next(o for o in model['slides'][0]['shapes'] if o.get('text') and o['text']['paragraphs'][0]['text']=='First')
        self.assertEqual(body['text']['paragraphs'][0]['bulletFont'],'Wingdings')
        self.assertEqual(body['text']['paragraphs'][0]['bullet'],'§')
        pp.remove(pp.find('{'+A+'}buFont'));pp.append(OxmlElement('a:buFontTx'))
        stream=BytesIO();prs.save(stream);model,_=parse_pptx(stream.getvalue())
        body=next(o for o in model['slides'][0]['shapes'] if o.get('text') and o['text']['paragraphs'][0]['text']=='First')
        self.assertIsNone(body['text']['paragraphs'][0]['bulletFont'])


if __name__=='__main__': unittest.main()
