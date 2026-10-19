"""Rendering regressions for native OOXML point spacing and kerning."""
import io
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx_parser import parse_pptx

class LineMetricsTests(unittest.TestCase):
    def test_absolute_point_line_spacing_is_exposed_in_pixels(self):
        prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3))
        para=box.text_frame.paragraphs[0]
        para.text='First line\nSecond line'
        para.line_spacing=Pt(22.5)
        out=io.BytesIO();prs.save(out)
        model,_=parse_pptx(out.getvalue())
        text=next(obj['text'] for obj in model['slides'][0]['shapes']
                  if obj['source']=='slide' and obj.get('text'))
        p=text['paragraphs'][0]
        self.assertAlmostEqual(p['lineSpacingPx'],30.0,places=2)
        self.assertIsNone(p['lineSpacing'])

    def test_ratio_line_spacing_is_retained_as_ratio(self):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3))
        para=box.text_frame.paragraphs[0];para.text='A B C';para.line_spacing=1.2
        out=io.BytesIO();prs.save(out)
        model,_=parse_pptx(out.getvalue())
        text=next(obj['text'] for obj in model['slides'][0]['shapes']
                  if obj['source']=='slide' and obj.get('text'))
        p=text['paragraphs'][0]
        self.assertAlmostEqual(p['lineSpacing'],1.2,places=4)
        self.assertIsNone(p['lineSpacingPx'])

    def test_browser_kerning_respects_run_override_and_smartart(self):
        js=(Path(__file__).resolve().parents[1]/'static/viewer.js').read_text()
        self.assertIn("if (!opts.smartArt) box.style.fontKerning = 'none'",js)
        self.assertIn('opts.smartArt',js)
        self.assertIn('run.kerning != null',js)

if __name__=='__main__':unittest.main()
