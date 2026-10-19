"""Creates deterministic sample deck; developer-only, not needed to run viewer."""
from pathlib import Path
from io import BytesIO
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData

ROOT=Path(__file__).resolve().parents[1]
prs=Presentation()
prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
NAVY=(21,34,46);DARK=(13,23,32);TEAL=(87,204,186);WHITE=(244,248,251);MUTED=(155,178,192);GOLD=(238,184,99)

def rgb(t):return RGBColor(*t)
def shape(slide,kind,x,y,w,h,fill=None,line=None,width=1):
 s=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
 if fill is None: s.fill.background()
 else: s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
 if line is None: s.line.fill.background()
 else: s.line.color.rgb=rgb(line);s.line.width=Pt(width)
 return s

def txt(slide,x,y,w,h,text,size=20,color=WHITE,bold=False,font='Aptos',align=None,margin=0):
 s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=s.text_frame;tf.clear();tf.margin_left=Inches(margin);tf.margin_top=Inches(margin);tf.margin_right=Inches(margin);tf.margin_bottom=Inches(margin)
 p=tf.paragraphs[0];p.text=text;p.alignment=align
 for r in p.runs: r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
 return s

def bg(slide,color=DARK):
 slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(color)

def footer(slide,n,label):
 shape(slide,MSO_SHAPE.RECTANGLE,.55,7.11,12.25,.014,fill=(58,76,87))
 txt(slide,.56,7.18,4,.2,'S L I D E S C O P E   /   RENDERER QA',8,MUTED,bold=True)
 txt(slide,12.12,7.16,.65,.22,f'{n:02d} / 04',8,TEAL,bold=True)

# slide 1
s=prs.slides.add_slide(blank);bg(s)
shape(s,MSO_SHAPE.RECTANGLE,.62,.71,.12,.62,TEAL)
txt(s,.91,.7,7.4,.3,'RENDER ENGINE  /  VISUAL TEST DECK',11,TEAL,bold=True)
txt(s,.62,1.57,10.2,1.04,'Make every slide',47,WHITE,bold=True)
txt(s,.62,2.48,11,1.0,'inspectable.',50,TEAL,bold=True)
txt(s,.65,3.74,7.25,.97,'A browser-native PowerPoint preview that gives AI agents a clear window into the presentation itself.',19,MUTED)
shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,.66,5.2,3.45,.64,fill=(39,69,73))
txt(s,.84,5.36,3.05,.3,'OFFLINE  •  PYTHON 3.12',13,TEAL,bold=True)
for i,(x,y,r,c) in enumerate([(10.37,1.45,1.1,(37,80,85)),(11.31,2.25,.88,TEAL),(9.49,3.38,1.28,(54,109,109)),(11.25,4.52,.43,GOLD)]):
 shape(s,MSO_SHAPE.OVAL,x,y,r,r,fill=c)
footer(s,1,'Cover')

# slide 2
s=prs.slides.add_slide(blank);bg(s,(245,248,249))
txt(s,.62,.42,6,.35,'02 / COMPONENT COVERAGE',12,(39,135,128),True)
txt(s,.62,.9,9,.8,'Visual primitives & layout',32,(28,47,58),True)
txt(s,.63,1.7,9,.38,'Individual elements should render with accurate bounds and editable text runs.',16,(102,122,134))
metrics=[('04','LAYOUT PRIMITIVES'),('09','PRESET GEOMETRIES'),('100%','BROWSER-NATIVE')]
for i,(num,label) in enumerate(metrics):
 x=.64+i*4.15
 shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,2.43,3.88,1.48,fill=(255,255,255),line=(219,231,232))
 txt(s,x+.22,2.6,3.25,.56,num,30,(26,118,113),True)
 txt(s,x+.22,3.38,3.24,.27,label,10,(100,116,126),True)
txt(s,.64,4.33,5.3,.5,'Shape and color accuracy',21,(28,47,58),True)
shape(s,MSO_SHAPE.CHEVRON,.76,5.05,1.06,.99,TEAL)
shape(s,MSO_SHAPE.DIAMOND,2.48,5.02,.99,.99,GOLD)
shape(s,MSO_SHAPE.HEXAGON,4.12,5.03,1.08,1.0,(103,135,200))
shape(s,MSO_SHAPE.OVAL,6.11,5.01,1.06,1.04,(126,172,173))
shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,8,5,3.88,1.10,fill=(38,73,83))
txt(s,8.28,5.34,3.2,.42,'ROUNDED RECTANGLE',13,WHITE,True)
footer(s,2,'Shapes')

# slide 3
s=prs.slides.add_slide(blank);bg(s,(248,250,250))
txt(s,.62,.45,9,.4,'03 / DATA & CHART RENDERING',12,(36,130,124),True)
txt(s,.62,.96,8,.65,'A chart should stay readable.',31,(33,54,67),True)
txt(s,.62,1.67,9,.37,'A basic data chart, a structured table, and small annotations.',16,(107,124,135))
chart_data=CategoryChartData();chart_data.categories=['Q1','Q2','Q3','Q4'];chart_data.add_series('North',[12,19,24,29]);chart_data.add_series('South',[9,16,21,27]);chart_data.add_series('West',[6,11,16,23])
chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.68),Inches(2.35),Inches(7.25),Inches(3.98),chart_data).chart
chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM
chart.has_title=True;chart.chart_title.text_frame.text='Region performance (sample)'
for series in chart.series:
 try:series.format.fill.solid();series.format.fill.fore_color.rgb=rgb((54,125,143))
 except Exception:pass
rows,cols=5,2
frame=s.shapes.add_table(rows,cols,Inches(8.5),Inches(2.44),Inches(4.23),Inches(3.72))
table=frame.table;table.columns[0].width=Inches(2.2);table.columns[1].width=Inches(2.03)
items=[('METRIC','VALUE'),('Text boxes','Pass'),('Shapes','Pass'),('Charts','Preview'),('Grouped objects','Preview')]
for r,(a,b) in enumerate(items):
 for c,val in enumerate((a,b)):
  cell=table.cell(r,c);cell.text=val;cell.fill.solid();cell.fill.fore_color.rgb=rgb((34,73,81) if r==0 else ((232,241,241) if r%2 else (248,251,251)))
  for p in cell.text_frame.paragraphs:
   for run in p.runs:
    run.font.name='Aptos';run.font.size=Pt(12 if r else 11);run.font.bold=r==0;run.font.color.rgb=rgb(WHITE if r==0 else (39,68,79))
footer(s,3,'Data')

# slide 4
s=prs.slides.add_slide(blank);bg(s,(24,38,51))
txt(s,.62,.42,9,.37,'04 / IMAGES + RICH TEXT',12,TEAL,True)
txt(s,.62,1.03,9,.62,'More than a screenshot',32,WHITE,True)
txt(s,.63,1.8,7,.41,'Objects can be selected, measured and enumerated.',17,MUTED)
img=Image.new('RGB',(1000,650),(19,50,67)); d=ImageDraw.Draw(img)
for i in range(44):
 x=i*28-180
 d.polygon([(x,650),(x+500,0),(x+575,0),(x+75,650)],fill=(30+i*2%65,85+i%58,99+i%62))
for k in range(7):
 yy=80+k*72
 d.line([(50,yy),(940,yy)],fill=(103,205,183),width=3)
buf=BytesIO();img.save(buf,format='PNG');buf.seek(0)
s.shapes.add_picture(buf,Inches(.68),Inches(2.57),width=Inches(6.0),height=Inches(3.80))
shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,7.07,2.56,5.48,3.83,(33,55,66))
txt(s,7.45,2.91,4.5,.35,'INSPECTION READY',14,TEAL,True)
txt(s,7.45,3.47,4.4,.6,'Select any object.',23,WHITE,True)
txt(s,7.45,4.14,4.5,1.22,'The browser exposes slide objects with stable IDs, geometry, text content, and formatting properties.',16,MUTED)
shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,7.45,5.68,3.35,.42,(53,117,117))
txt(s,7.61,5.77,2.9,.22,'PLAYWRIGHT FRIENDLY',11,WHITE,True)
footer(s,4,'Objects')

out=ROOT/'sample'/'demo.pptx';out.parent.mkdir(exist_ok=True)
prs.save(out)
print('Created',out,out.stat().st_size)
