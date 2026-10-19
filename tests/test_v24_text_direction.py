"""Text properties which must survive parsing into inspectable browser runs."""
import unittest
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement
from pptx_parser import parse_pptx


class TextDirectionTests(unittest.TestCase):
    def model(self, direction='horz', rotation=0, baseline=None, default_baseline=None):
        prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(1),Inches(3))
        shape.text_frame.word_wrap=False
        body=shape.text_frame._txBody.bodyPr
        body.set('vert',direction);body.set('rot',str(rotation))
        para=shape.text_frame.paragraphs[0]
        if default_baseline is not None:
            props=para._p.get_or_add_pPr();defaults=OxmlElement('a:defRPr')
            defaults.set('baseline',str(default_baseline));props.append(defaults)
        run=para.add_run();run.text='Label';run.font.size=Pt(20)
        if baseline is not None:run._r.get_or_add_rPr().set('baseline',str(baseline))
        out=BytesIO();prs.save(out);model,_=parse_pptx(out.getvalue())
        return next(s['text'] for s in model['slides'][0]['shapes'] if s['name']==shape.name)

    def test_vertical_modes_survive_without_changing_anchor(self):
        for mode in ('vert','vert270','eaVert'):
            result=self.model(direction=mode)
            self.assertEqual(result['direction'],mode)
            self.assertEqual(result['vertical'],'top')
            self.assertFalse(result['wrap'])

    def test_text_rotation_is_separate_from_shape_rotation(self):
        self.assertEqual(self.model(rotation=-5400000)['textRotation'],-90)

    def test_script_offset_preserves_declared_font_size(self):
        for value in (-25000,30000):
            run=self.model(baseline=value)['paragraphs'][0]['runs'][0]
            self.assertEqual(run['baseline'],value/100000)
            self.assertAlmostEqual(run['size'],26.667,places=3)

    def test_paragraph_default_and_explicit_baseline_reset(self):
        inherited=self.model(default_baseline=30000)['paragraphs'][0]['runs'][0]
        reset=self.model(default_baseline=30000,baseline=0)['paragraphs'][0]['runs'][0]
        self.assertEqual(inherited['baseline'],.3)
        self.assertEqual(reset['baseline'],0)


if __name__=='__main__':unittest.main()
