"""Self-contained regression cases introduced by the 59-slide SLS reference deck.
No corporate files or extra packages are required to run these tests.
"""
import unittest
from io import BytesIO
from pathlib import Path
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from pptx_parser import parse_pptx

A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def rgb_image():
    blob=BytesIO()
    Image.new('RGB',(16,12),(60,120,200)).save(blob,format='PNG')
    blob.seek(0)
    return blob


def deck():
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    picture=slide.shapes.add_picture(rgb_image(),Inches(1),Inches(1),Inches(2),Inches(2))
    blip=picture._element.xpath('.//a:blip')[0]
    # A master-effect-like image brightness/contrast transform.
    etree.SubElement(blip,A+'lum',bright='-60000',contrast='-63000')
    # A picture embedded in an ordinary shape's p:spPr instead of a p:pic.
    rectangle=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4),Inches(1), Inches(2),Inches(1))
    sppr=rectangle._element.spPr
    for child in list(sppr):
        if etree.QName(child).localname in ('solidFill','gradFill','noFill'):
            sppr.remove(child)
    fill=etree.SubElement(sppr,A+'blipFill')
    etree.SubElement(fill,A+'blip', {R+'embed':blip.get(R+'embed')})
    src=etree.SubElement(fill,A+'srcRect')
    src.set('l','10000')
    src.set('r','15000')
    stretch=etree.SubElement(fill,A+'stretch')
    etree.SubElement(stretch,A+'fillRect')
    slide.shapes.add_shape(MSO_SHAPE.FRAME, Inches(4),Inches(3), Inches(2), Inches(1))
    data=CategoryChartData()
    data.categories=['A','B','C']
    data.add_series('Values',(1.0,2.0,3.0))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(4),Inches(4), Inches(2),data)
    stream=BytesIO();prs.save(stream);return stream.getvalue()


class TestV20Fidelity(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.model,cls.assets=parse_pptx(deck(),'synthetic.pptx')
        cls.shapes=cls.model['slides'][0]['shapes']

    def test_luminance_is_read_from_a_blip(self):
        image=next(s for s in self.shapes if s['kind']=='image')
        self.assertEqual(image['image']['luminance'],{'brightness':-.6,'contrast':-.63})

    def test_picture_filled_rectangle_is_not_dropped(self):
        shape=next(s for s in self.shapes if s.get('fillImage'))
        self.assertEqual(shape['kind'],'shape')
        self.assertEqual(shape['fillImage']['crop'],[.1,0,.15,0])
        self.assertTrue(shape['fillImage']['url'].startswith('/api/media/'))

    def test_picture_filled_shape_reuses_embedded_asset(self):
        image=next(s for s in self.shapes if s['kind']=='image')
        shape=next(s for s in self.shapes if s.get('fillImage'))
        self.assertEqual(image['image']['url'],shape['fillImage']['url'])

    def test_hollow_frame_preset_is_distinguished(self):
        frame=next(s for s in self.shapes if s['preset']=='frame')
        self.assertGreater(frame['w'],0)
        self.assertGreater(frame['frameAdj'],0)
        self.assertLess(frame['frameAdj'],.5)

    def test_native_chart_preserves_transparency(self):
        chart=next(s for s in self.shapes if s['kind']=='chart')
        self.assertTrue(chart['chart']['series'])
        self.assertEqual(chart['chart']['background'],'transparent')

    def test_svg_chart_labels_are_created_as_nodes(self):
        js=(Path(__file__).resolve().parents[1]/'static/viewer.js').read_text()
        self.assertNotIn(')).lastChild.textContent',js)
        self.assertIn('const label=svgEl(\'text\'',js)

    def test_frame_is_hollow_in_browser(self):
        js=(Path(__file__).resolve().parents[1]/'static/viewer.js').read_text()
        self.assertIn("if (kind === 'frame')",js)
        self.assertIn("'fill-rule':'evenodd'",js)


if __name__=='__main__':
    unittest.main()
