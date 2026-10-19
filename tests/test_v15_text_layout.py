"""OOXML typography regression cases for the browser renderer (v1.5)."""
from __future__ import annotations
import io
import os
from pathlib import Path
import unittest

from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches
from pptx_parser import parse_pptx


class TextLayoutTests(unittest.TestCase):
    def test_direct_run_kerning_is_exposed_to_browser_model(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = 'No pair kerning: AV WA To'
        rpr = run._r.get_or_add_rPr()
        rpr.set('kern', '0')
        bio = io.BytesIO()
        prs.save(bio)
        data, _ = parse_pptx(bio.getvalue())
        para = next(sh['text']['paragraphs'][0] for sh in data['slides'][0]['shapes']
                    if sh['source'] == 'slide' and sh.get('text') and
                    sh['text']['paragraphs'][0]['text'].startswith('No pair'))
        self.assertEqual(para['runs'][0]['kerning'], '0')

    def test_ordinary_runs_do_not_disable_kerning(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        shape.text = 'Default kerning'
        bio = io.BytesIO()
        prs.save(bio)
        model, _ = parse_pptx(bio.getvalue())
        para = next(sh['text']['paragraphs'][0] for sh in model['slides'][0]['shapes']
                    if sh['source'] == 'slide' and sh.get('text') and
                    sh['text']['paragraphs'][0]['text'] == 'Default kerning')
        self.assertIsNone(para['runs'][0]['kerning'])

    def test_reference_deck_contains_scaled_blanks_and_subtitles(self):
        name = os.environ.get('SLIDESCOPE_REGRESSION_PPTX')
        if not name:
            self.skipTest('Reference presentation is optional')
        model, _ = parse_pptx(Path(name).read_bytes())
        self.assertEqual(model['count'], 12)
        for index in (0, 8):
            subtitles = [sh for sh in model['slides'][index]['shapes']
                         if sh.get('placeholderType') == 'SUBTITLE' and sh.get('text')]
            self.assertEqual(len(subtitles), 1)
            self.assertEqual(len(subtitles[0]['text']['paragraphs']), 3)
            self.assertTrue(all(p['lineSpacing'] is None
                                for p in subtitles[0]['text']['paragraphs']))
        body = next(sh for sh in model['slides'][1]['shapes']
                    if sh.get('name') == 'Content Placeholder 4' and sh['source'] == 'slide')
        self.assertEqual(body['text']['autoSize'], 'TEXT_TO_FIT_SHAPE')
        self.assertLess(body['text']['fontScale'], 1)
        self.assertEqual([i for i, p in enumerate(body['text']['paragraphs'])
                          if not p['text']], [1, 3])
        # This presenter disabled PowerPoint pair kerning in some slide-6 runs.
        self.assertTrue(any(run.get('kerning') == '0'
                            for sh in model['slides'][5]['shapes'] if sh.get('text')
                            for para in sh['text']['paragraphs'] for run in para['runs']))


if __name__ == '__main__':
    unittest.main()
